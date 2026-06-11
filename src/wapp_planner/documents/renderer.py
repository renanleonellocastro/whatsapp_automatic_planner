"""Render layout blocks to PDF and DOCX bytes (FR-RND-01/02).

Both backends consume the same :func:`build_blocks` output so the PDF and Word
documents stay consistent. :func:`render_document` returns ready-to-send
attachments for both formats (AD-4).
"""

from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import ListFlowable, ListItem, SimpleDocTemplate, Spacer, TableStyle
from reportlab.platypus import Paragraph as PdfParagraph
from reportlab.platypus import Table as PdfTable

from wapp_planner.documents.branding import Branding
from wapp_planner.documents.content import ExecutionPlanContent, QuoteContent
from wapp_planner.documents.layout import (
    Block,
    Bullets,
    Heading,
    Paragraph,
    Table,
    build_blocks,
)
from wapp_planner.providers.base import Attachment

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def render_pdf(blocks: list[Block]) -> bytes:
    """Render blocks to a PDF document."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, title="document")
    styles = getSampleStyleSheet()
    story: list[object] = []
    for block in blocks:
        if isinstance(block, Heading):
            style = styles["Heading1"] if block.level == 1 else styles["Heading2"]
            story.append(PdfParagraph(escape(block.text), style))
        elif isinstance(block, Paragraph):
            story.append(PdfParagraph(escape(block.text), styles["BodyText"]))
        elif isinstance(block, Bullets):
            story.append(PdfParagraph(escape(block.title), styles["Heading3"]))
            story.append(
                ListFlowable(
                    [
                        ListItem(PdfParagraph(escape(item), styles["BodyText"]))
                        for item in block.items
                    ],
                    bulletType="bullet",
                )
            )
        elif isinstance(block, Table):
            table = PdfTable([block.headers, *block.rows])
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, "#888888")]))
            story.append(table)
        else:  # KeyValues
            for key, value in block.pairs:
                story.append(
                    PdfParagraph(f"<b>{escape(key)}:</b> {escape(value)}", styles["BodyText"])
                )
        story.append(Spacer(1, 8))
    doc.build(story)
    return buffer.getvalue()


def render_docx(blocks: list[Block]) -> bytes:
    """Render blocks to a Word (.docx) document."""
    doc = DocxDocument()
    for block in blocks:
        if isinstance(block, Heading):
            doc.add_heading(block.text, level=block.level)
        elif isinstance(block, Paragraph):
            doc.add_paragraph(block.text)
        elif isinstance(block, Bullets):
            doc.add_heading(block.title, level=3)
            for item in block.items:
                doc.add_paragraph(item, style="List Bullet")
        elif isinstance(block, Table):
            table = doc.add_table(rows=1 + len(block.rows), cols=len(block.headers))
            table.style = "Table Grid"
            for col, header in enumerate(block.headers):
                table.rows[0].cells[col].text = header
            for row_index, row in enumerate(block.rows, start=1):
                for col, cell_value in enumerate(row):
                    table.rows[row_index].cells[col].text = cell_value
        else:  # KeyValues
            for key, value in block.pairs:
                paragraph = doc.add_paragraph()
                paragraph.add_run(f"{key}: ").bold = True
                paragraph.add_run(value)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_document(
    content: QuoteContent | ExecutionPlanContent, branding: Branding, *, basename: str
) -> list[Attachment]:
    """Render both PDF and DOCX attachments from document content (AD-4)."""
    blocks = build_blocks(content, branding)
    return [
        Attachment(filename=f"{basename}.pdf", content=render_pdf(blocks), content_type=PDF_MIME),
        Attachment(
            filename=f"{basename}.docx", content=render_docx(blocks), content_type=DOCX_MIME
        ),
    ]
