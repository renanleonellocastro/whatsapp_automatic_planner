"""Unit tests for PDF/DOCX rendering (FR-RND-01/02, AD-4)."""

from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument

from wapp_planner.documents.branding import Branding
from wapp_planner.documents.content import LineItem, QuoteContent
from wapp_planner.documents.layout import Bullets, Heading, KeyValues, Paragraph, Table
from wapp_planner.documents.renderer import (
    DOCX_MIME,
    PDF_MIME,
    render_document,
    render_docx,
    render_pdf,
)

ALL_BLOCKS = [
    Heading("Title", 1),
    Heading("Section", 2),
    Paragraph("Some body text with an & ampersand"),
    Bullets("Steps", ["one", "two"]),
    Table(["A", "B"], [["1", "2"], ["3", "4"]]),
    KeyValues([("Total", "$10.00")]),
]


def test_render_pdf_is_a_valid_pdf() -> None:
    data = render_pdf(ALL_BLOCKS)
    assert data.startswith(b"%PDF")
    assert len(data) > 500


def test_render_docx_is_openable_and_has_text() -> None:
    data = render_docx(ALL_BLOCKS)
    assert data[:2] == b"PK"  # zip/OOXML signature
    doc = DocxDocument(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Title" in text
    assert "one" in text
    assert "Total: " in text
    # the table rendered with header + 2 rows
    assert doc.tables[0].rows[0].cells[0].text == "A"
    assert doc.tables[0].rows[2].cells[1].text == "4"


def _quote() -> QuoteContent:
    return QuoteContent(
        title="Estimate",
        prepared_for="Customer",
        scope_summary="scope",
        line_items=[
            LineItem(description="Demo", quantity=1, unit="job", unit_price=500, total=500)
        ],
        subtotal=500,
        total=500,
        assumptions=["a"],
    )


def test_render_document_returns_pdf_and_docx_attachments() -> None:
    attachments = render_document(_quote(), Branding(company_name="Acme"), basename="quote-req-1")
    assert [a.filename for a in attachments] == ["quote-req-1.pdf", "quote-req-1.docx"]
    pdf, docx = attachments
    assert pdf.content_type == PDF_MIME
    assert pdf.content.startswith(b"%PDF")
    assert docx.content_type == DOCX_MIME
    assert docx.content[:2] == b"PK"
